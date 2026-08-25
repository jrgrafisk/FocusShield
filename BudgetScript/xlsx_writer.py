# -*- coding: utf-8 -*-
"""Write the budget into a plain .xlsx file (no macros).

Mirrors LibreOfficeBudget/office/budget_office.py in spirit, but targets
openpyxl instead of UNO, and is simpler in places (no hover comments with
transaction samples, no per-account forecast breakdown, no Bankbudget
snapshot in v1). All the actual parsing/categorisation/classification
logic is reused as-is from budget_core - this module only writes cells.
"""

from __future__ import annotations

import datetime as _dt
from typing import Dict, List, Optional, Sequence, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from budget_core.budget import KIND_EXPENSE, KIND_INCOME, Summary, Transaction
from budget_core.plan import (GROUP_FIXED, GROUP_LABELS, GROUP_PERIODIC,
                              GROUP_VARIABLE, BudgetPlan, build_plan)
from budget_core.rules import DEFAULT_CATEGORY, IGNORED_CATEGORY, RuleSet
from budget_core.textutils import fold
from budget_core.transfers import parse_account_numbers, text_mentions_account

__all__ = ["BudgetBook"]

# ---------------------------------------------------------------------------
# Design tokens (same palette as the LibreOffice/Excel versions)
# ---------------------------------------------------------------------------
NAVY = "334960"
ORANGE = "F46524"
TEXT_GREY = "576475"
MUTED = "687887"
DARK = "434343"
PEACH = "FFF2ED"
LIGHT_TEXT = "CCCCCC"

FONT_BODY = "Calibri"
FONT_TITLE = "Calibri Light"

FMT_CURRENCY = '#,##0 "kr."'
FMT_CURRENCY2 = '#,##0.00 "kr."'
FMT_SIGNED = '+#,##0 "kr.";-#,##0 "kr."'
FMT_DATE = "dd-mm-yyyy"
FMT_MONTH = "mmmm yyyy"

SHEET_TX = "Transaktioner"
SHEET_RULES = "Kategorier"
SHEET_PLAN = "Budgetforslag"
SHEET_FORECAST = "Prognose"
SHEET_ACCOUNTS = "Konti"
SHEET_MONTHS = "Alle måneder"

TX_HEADER_ROW = 4
TX_FIRST_ROW = 5
TX_MAX_ROW = 5000
COL_EXP_DATE, COL_EXP_AMOUNT, COL_EXP_TEXT, COL_EXP_CAT = 2, 3, 4, 5
COL_INC_DATE, COL_INC_AMOUNT, COL_INC_TEXT, COL_INC_CAT = 7, 8, 9, 10

ACC_HEADER_ROW = 4
ACC_FIRST_ROW = 5
ACC_COL_NAME, ACC_COL_TYPE, ACC_COL_BALANCE, ACC_COL_NUMBER = 2, 3, 4, 5

PLAN_HEADER_ROW = 5
FORECAST_MONTHS = 24
FORECAST_HEADER_ROW = 9


def _hex(color: str) -> str:
    return "FF" + color


class Pen:
    """Small styling helper wrapping one worksheet."""

    def __init__(self, ws: Worksheet):
        self.ws = ws

    def cell(self, ref: str):
        return self.ws[ref]

    def text(self, ref, value, *, bold=False, size=10, color=DARK, align="left",
            italic=False, wrap=False, valign=None, bg=None, font=FONT_BODY):
        cell = self.ws[ref]
        cell.value = value
        self._style(cell, bold, size, color, align, italic, wrap, valign, bg, font)
        return cell

    def number(self, ref, value, *, fmt=FMT_CURRENCY, bold=False, size=10,
              color=DARK, align="right", bg=None, font=FONT_BODY):
        cell = self.ws[ref]
        cell.value = value
        cell.number_format = fmt
        self._style(cell, bold, size, color, align, False, False, None, bg, font)
        return cell

    def date(self, ref, value, *, fmt=FMT_DATE, bold=False, size=10, color=DARK,
            align="left", font=FONT_BODY):
        cell = self.ws[ref]
        cell.value = value
        cell.number_format = fmt
        self._style(cell, bold, size, color, align, False, False, None, None, font)
        return cell

    def formula(self, ref, formula_text, *, fmt=FMT_CURRENCY, bold=False, size=10,
               color=DARK, align="right", bg=None, font=FONT_BODY):
        cell = self.ws[ref]
        cell.value = formula_text
        if fmt:
            cell.number_format = fmt
        self._style(cell, bold, size, color, align, False, False, None, bg, font)
        return cell

    def style(self, ref, *, color=None, italic=None, bold=None, fmt=None, align=None, bg=None):
        for row in self.ws[ref]:
            for cell in row:
                current = cell.font
                cell.font = Font(
                    name=current.name or FONT_BODY,
                    size=current.size or 10,
                    bold=bold if bold is not None else current.bold,
                    italic=italic if italic is not None else current.italic,
                    color=_hex(color) if color is not None else current.color)
                if fmt:
                    cell.number_format = fmt
                if align:
                    cell.alignment = Alignment(horizontal=align, vertical=cell.alignment.vertical)
                if bg is not None:
                    cell.fill = PatternFill("solid", fgColor=_hex(bg))

    def _style(self, cell, bold, size, color, align, italic, wrap, valign, bg, font):
        cell.font = Font(name=font, size=size, bold=bold, italic=italic,
                         color=_hex(color) if color else None)
        cell.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=wrap)
        if bg is not None:
            cell.fill = PatternFill("solid", fgColor=_hex(bg))

    def merge(self, ref):
        self.ws.merge_cells(ref)

    def column_width(self, col_letter, width):
        self.ws.column_dimensions[col_letter].width = width

    def row_height(self, row, height):
        self.ws.row_dimensions[row].height = height

    def note(self, ref, text):
        if not text:
            return
        self.ws[ref].comment = Comment(text, "Budget fra CSV", width=260, height=120)

    def dropdown(self, ref, formula):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        self.ws.add_data_validation(dv)
        dv.add(ref)

    def freeze(self, ref):
        self.ws.freeze_panes = ref


def _col(index: int) -> str:
    return get_column_letter(index)


def _ref(col: int, row: int) -> str:
    return f"{_col(col)}{row}"


def _range(col1: int, row1: int, col2: int, row2: int) -> str:
    return f"{_ref(col1, row1)}:{_ref(col2, row2)}"


def _next_month(month_key: str) -> _dt.date:
    if not month_key:
        today = _dt.date.today()
        return _dt.date(today.year, today.month, 1)
    year, month = (int(p) for p in month_key.split("-"))
    return _add_months(_dt.date(year, month, 1), 1)


def _add_months(date: _dt.date, count: int) -> _dt.date:
    total = date.month - 1 + count
    year = date.year + total // 12
    month = total % 12 + 1
    return _dt.date(year, month, 1)


def _month_start(month_key: str) -> _dt.date:
    year, month = (int(p) for p in month_key.split("-"))
    return _dt.date(year, month, 1)


class SheetTransaction:
    """A transaction read back from the workbook."""

    __slots__ = ("date", "text", "amount", "category", "source_row", "income")

    def __init__(self, date, text, amount, category, source_row, income):
        self.date = date
        self.text = text
        self.amount = amount
        self.category = category
        self.source_row = source_row
        self.income = income

    @property
    def month(self) -> str:
        return "%04d-%02d" % (self.date.year, self.date.month)

    @property
    def kind(self) -> str:
        return KIND_INCOME if self.income else KIND_EXPENSE


class BudgetBook:
    """Reads and writes one budget workbook (.xlsx, no macros)."""

    def __init__(self, wb: Optional[Workbook] = None):
        self.wb = wb or Workbook()

    # -- sheet helpers ------------------------------------------------
    def sheet(self, name: str) -> Worksheet:
        if name in self.wb.sheetnames:
            del self.wb[name]
        return self.wb.create_sheet(name)

    def has_sheet(self, name: str) -> bool:
        return name in self.wb.sheetnames

    def pen(self, ws: Worksheet) -> Pen:
        return Pen(ws)

    def used_row_count(self, ws: Worksheet) -> int:
        return ws.max_row or 1

    # -- Transaktioner --------------------------------------------------
    def write_transactions(self, ws: Worksheet, transactions: Sequence[Transaction]) -> None:
        pen = self.pen(ws)
        widths = {1: 4, 2: 11, 3: 15, 4: 30, 5: 16, 6: 4, 7: 11, 8: 15, 9: 30, 10: 16}
        for col, w in widths.items():
            pen.column_width(_col(col), w)

        pen.merge("B1:J1")
        pen.text("B1", "Skift eller tilføj kategorier i kolonnerne nedenfor - og i "
                       "arket \"Kategorier\", hvis de skal sættes automatisk.",
                 size=10, color=LIGHT_TEXT, bg=NAVY, italic=True, align="left", valign="center")
        pen.row_height(1, 22)

        pen.text("B2", "Udgifter", font=FONT_TITLE, size=18, bold=True, color=ORANGE)
        pen.text("G2", "Indtægter", font=FONT_TITLE, size=18, bold=True, color=ORANGE)
        pen.row_height(2, 26)

        headers = ("Dato", "Beløb", "Beskrivelse", "Kategori")
        for offset, title in enumerate(headers):
            pen.text(_ref(COL_EXP_DATE + offset, TX_HEADER_ROW), title, bold=True, size=11, color=NAVY)
            pen.text(_ref(COL_INC_DATE + offset, TX_HEADER_ROW), title, bold=True, size=11, color=NAVY)

        expenses = [t for t in transactions if t.amount < 0]
        income = [t for t in transactions if t.amount >= 0]
        self._write_block(ws, expenses, COL_EXP_DATE, TX_FIRST_ROW)
        self._write_block(ws, income, COL_INC_DATE, TX_FIRST_ROW)

        last = TX_FIRST_ROW + max(len(expenses), len(income), 1) - 1
        self._style_tx_range(pen, last)
        self._dim_ignored_rows(pen, expenses, COL_EXP_DATE, TX_FIRST_ROW)
        self._dim_ignored_rows(pen, income, COL_INC_DATE, TX_FIRST_ROW)
        self._add_category_dropdown(ws, COL_EXP_CAT, last)
        self._add_category_dropdown(ws, COL_INC_CAT, last)
        pen.freeze(_ref(2, TX_FIRST_ROW))

    def _style_tx_range(self, pen: Pen, last: int) -> None:
        for first_col in (COL_EXP_DATE, COL_INC_DATE):
            pen.style(_range(first_col, TX_FIRST_ROW, first_col, last), fmt=FMT_DATE,
                     color=MUTED, align="left")
            pen.style(_range(first_col + 1, TX_FIRST_ROW, first_col + 1, last), fmt=FMT_CURRENCY2,
                     color=TEXT_GREY, bold=True, align="left")
            pen.style(_range(first_col + 2, TX_FIRST_ROW, first_col + 3, last), color=TEXT_GREY,
                     align="left")

    def _write_block(self, ws: Worksheet, transactions: Sequence[Transaction], first_col: int,
                     start_row: int) -> None:
        for offset, t in enumerate(transactions):
            row = start_row + offset
            ws.cell(row, first_col, t.date)
            ws.cell(row, first_col + 1, abs(t.amount))
            ws.cell(row, first_col + 2, t.text)
            ws.cell(row, first_col + 3, t.category)

    def _dim_ignored_rows(self, pen: Pen, transactions: Sequence[Transaction], first_col: int,
                          start_row: int) -> None:
        for offset, t in enumerate(transactions):
            if t.category == IGNORED_CATEGORY:
                row = start_row + offset
                pen.style(_range(first_col, row, first_col + 3, row), color=MUTED, italic=True)

    def _add_category_dropdown(self, ws: Worksheet, col: int, last_row: int) -> None:
        if last_row < TX_FIRST_ROW:
            return
        pen = self.pen(ws)
        pen.dropdown(_range(col, TX_FIRST_ROW, col, last_row), f"'{SHEET_RULES}'!$E$5:$E$80")

    def read_transactions(self) -> List[SheetTransaction]:
        if not self.has_sheet(SHEET_TX):
            return []
        ws = self.wb[SHEET_TX]
        last_row = self.used_row_count(ws)
        if last_row < TX_FIRST_ROW:
            return []
        out: List[SheetTransaction] = []
        blocks = ((COL_EXP_DATE, COL_EXP_AMOUNT, COL_EXP_TEXT, COL_EXP_CAT, -1.0, False),
                  (COL_INC_DATE, COL_INC_AMOUNT, COL_INC_TEXT, COL_INC_CAT, 1.0, True))
        for date_col, amt_col, text_col, cat_col, sign, income in blocks:
            for row in range(TX_FIRST_ROW, last_row + 1):
                date_val = ws.cell(row, date_col).value
                if not isinstance(date_val, (_dt.date, _dt.datetime)):
                    continue
                amt_val = ws.cell(row, amt_col).value
                if not isinstance(amt_val, (int, float)) or amt_val == 0:
                    continue
                date = date_val.date() if isinstance(date_val, _dt.datetime) else date_val
                text = str(ws.cell(row, text_col).value or "")
                category = str(ws.cell(row, cat_col).value or "").strip()
                out.append(SheetTransaction(date, text, sign * abs(float(amt_val)),
                                            category, row, income))
        out.sort(key=lambda t: (t.date, t.source_row))
        return out

    def has_budget(self) -> bool:
        return self.has_sheet(SHEET_TX)

    # -- Kategorier -------------------------------------------------------
    def write_rules(self, ws: Worksheet, ruleset: RuleSet,
                    accounts: Optional[Dict[str, str]] = None) -> None:
        pen = self.pen(ws)
        for col, w in {1: 4, 2: 26, 3: 22, 4: 4, 5: 22, 6: 20}.items():
            pen.column_width(_col(col), w)

        pen.merge("B2:F2")
        pen.text("B2", "Kategorier", font=FONT_TITLE, size=18, bold=True, color=ORANGE, align="left")
        pen.merge("B3:F3")
        pen.text("B3", "Tilføj, ret eller slet regler i venstre tabel (nøgleord -> "
                       "kategori). Højre tabel viser alle kategorier og hvilken konto, "
                       "de trækkes fra.",
                 size=9, italic=True, color=MUTED, align="left", wrap=True, valign="center")
        pen.row_height(3, 34)

        pen.text("B4", "Nøgleord", bold=True, size=11, color=NAVY, align="left")
        pen.text("C4", "Kategori", bold=True, size=11, color=NAVY, align="left")
        pen.text("E4", "Kategori", bold=True, size=11, color=NAVY, align="left")
        pen.text("F4", "Konto", bold=True, size=11, color=NAVY, align="left")

        for r, (keyword, category) in enumerate(ruleset.to_rows(), start=1):
            pen.text(_ref(2, 4 + r), keyword, size=10, color=DARK, align="left")
            pen.text(_ref(3, 4 + r), category, size=10, color=DARK, align="left")

        accounts = accounts or {}
        categories = ruleset.categories()
        for c, category in enumerate(categories, start=1):
            pen.text(_ref(5, 4 + c), category, size=10, color=DARK, align="left")
            pen.text(_ref(6, 4 + c), accounts.get(category, ""), size=10, color=DARK, align="left")
        last_cat_row = 4 + max(len(categories), 1)
        pen.dropdown(_range(6, 5, 6, last_cat_row), f"'{SHEET_ACCOUNTS}'!$B$5:$B$50")

        pen.freeze(_ref(1, 5))

    def read_rules(self) -> List[Tuple[str, str]]:
        if not self.has_sheet(SHEET_RULES):
            return []
        ws = self.wb[SHEET_RULES]
        last_row = self.used_row_count(ws)
        out = []
        for row in range(5, last_row + 1):
            keyword = str(ws.cell(row, 2).value or "").strip()
            category = str(ws.cell(row, 3).value or "").strip()
            if keyword and category:
                out.append((keyword, category))
        return out

    def read_category_accounts(self) -> Dict[str, str]:
        if not self.has_sheet(SHEET_RULES):
            return {}
        ws = self.wb[SHEET_RULES]
        last_row = self.used_row_count(ws)
        out = {}
        for row in range(5, last_row + 1):
            category = str(ws.cell(row, 5).value or "").strip()
            account = str(ws.cell(row, 6).value or "").strip()
            if category and account:
                out[category] = account
        return out

    # -- Konti ------------------------------------------------------------
    def write_accounts(self, ws: Worksheet, rows: Sequence[Tuple[str, str, float, str]]) -> None:
        pen = self.pen(ws)
        for col, w in {1: 4, 2: 22, 3: 18, 4: 16, 5: 24}.items():
            pen.column_width(_col(col), w)

        pen.merge("B2:F2")
        pen.text("B2", "Konti", font=FONT_TITLE, size=18, bold=True, color=ORANGE, align="left")
        pen.merge("B3:F3")
        pen.text("B3", "Nuværende saldo - Prognosen bruger summen som startsaldo. Sæt "
                       "dit eget kontonummer under \"Kontonummer(e)\" og kør scriptet igen "
                       f"for at få overførsler mellem dine egne konti sat til \"{IGNORED_CATEGORY}\" "
                       "i stedet for at tælle som indtægt/udgift - adskil flere numre med komma.",
                 size=9, italic=True, color=MUTED, align="left", wrap=True, valign="center")
        pen.row_height(3, 60)

        pen.text("B4", "Konto", bold=True, size=11, color=NAVY, align="left")
        pen.text("C4", "Type", bold=True, size=11, color=NAVY, align="left")
        pen.text("D4", "Saldo", bold=True, size=11, color=NAVY, align="right")
        pen.text("E4", "Kontonummer(e)", bold=True, size=11, color=NAVY, align="left")

        rows = list(rows) or [("", "", 0.0, "")]
        for r, (name, kind, balance, number) in enumerate(rows, start=1):
            rr = ACC_HEADER_ROW + r
            pen.text(_ref(ACC_COL_NAME, rr), name, bold=True, size=10, color=DARK, align="left")
            pen.text(_ref(ACC_COL_TYPE, rr), kind, size=10, color=TEXT_GREY, align="left")
            pen.number(_ref(ACC_COL_BALANCE, rr), balance, fmt=FMT_CURRENCY, size=10,
                      color=DARK, align="right", bg=PEACH)
            pen.text(_ref(ACC_COL_NUMBER, rr), number, size=10, color=DARK, align="left", bg=PEACH)

        last_row = ACC_HEADER_ROW + len(rows)
        total_row = last_row + 2
        pen.text(_ref(ACC_COL_NAME, total_row), "I alt", bold=True, size=10, color=NAVY, align="left")
        pen.formula(_ref(ACC_COL_BALANCE, total_row),
                   f"=SUM({_range(ACC_COL_BALANCE, ACC_FIRST_ROW, ACC_COL_BALANCE, last_row)})",
                   fmt=FMT_CURRENCY, bold=True, color=NAVY, align="right")
        pen.freeze(_ref(1, ACC_HEADER_ROW + 1))

    def ensure_accounts_sheet(self) -> None:
        if self.has_sheet(SHEET_ACCOUNTS):
            return
        ws = self.sheet(SHEET_ACCOUNTS)
        self.write_accounts(ws, [
            ("Lønkonto", "Lønkonto", 0.0, ""),
            ("Budgetkonto", "Budgetkonto", 0.0, ""),
            ("Opsparingskonto", "Opsparingskonto", 0.0, ""),
        ])

    def scan_accounts(self) -> List[Tuple[str, str, float, str, int]]:
        if not self.has_sheet(SHEET_ACCOUNTS):
            return []
        ws = self.wb[SHEET_ACCOUNTS]
        last_row = self.used_row_count(ws)
        out = []
        for row in range(ACC_FIRST_ROW, last_row + 1):
            name = str(ws.cell(row, ACC_COL_NAME).value or "").strip()
            if not name or fold(name) == "i alt":
                continue
            kind = str(ws.cell(row, ACC_COL_TYPE).value or "").strip()
            balance_val = ws.cell(row, ACC_COL_BALANCE).value
            balance = float(balance_val) if isinstance(balance_val, (int, float)) else 0.0
            number = str(ws.cell(row, ACC_COL_NUMBER).value or "").strip()
            out.append((name, kind, balance, number, row))
        return out

    def read_account_numbers(self) -> List[str]:
        out: List[str] = []
        for _name, _kind, _balance, number, _row in self.scan_accounts():
            out.extend(parse_account_numbers(number))
        return out

    def accounts_total(self) -> float:
        return sum(balance for _n, _k, balance, _num, _r in self.scan_accounts())

    # -- Budgetforslag ------------------------------------------------------
    def write_plan(self, ws: Worksheet, plan: BudgetPlan,
                   targets: Optional[Dict[str, float]] = None,
                   accounts: Optional[Dict[str, str]] = None,
                   plan_accounts: Optional[Dict[str, str]] = None) -> Dict[str, object]:
        pen = self.pen(ws)
        for col, w in {1: 4, 2: 26, 3: 16, 4: 16, 5: 14, 6: 18}.items():
            pen.column_width(_col(col), w)

        pen.merge("B2:F2")
        pen.text("B2", "Budgetforslag", font=FONT_TITLE, size=18, bold=True, color=ORANGE, align="left")
        pen.merge("B3:F3")
        if plan.uncertain:
            intro = (f"Udkast til et fast månedsbudget. Bemærk: bygger kun på "
                    f"{len(plan.months)} måned(er) - tallene er usikre. "
                    "Ret Mål-kolonnen som du vil.")
        else:
            intro = ("Udkast til et fast månedsbudget ud fra hele perioden. Ret "
                     "Mål-kolonnen som du vil - resten genberegnes automatisk, når "
                     "du kører scriptet igen.")
        pen.text("B3", intro, size=9, italic=True, color=MUTED, align="left", wrap=True, valign="center")
        pen.row_height(3, 34)

        header = PLAN_HEADER_ROW
        pen.text(_ref(2, header), "Kategori", bold=True, size=11, color=NAVY, align="left")
        pen.text(_ref(3, header), "Gennemsnit/md", bold=True, size=10, color=NAVY, align="right")
        pen.text(_ref(4, header), "Mål", bold=True, size=11, color=NAVY, align="right")
        pen.text(_ref(5, header), "Forskel", bold=True, size=10, color=NAVY, align="right")
        pen.text(_ref(6, header), "Konto", bold=True, size=10, color=NAVY, align="left")

        targets = targets or {}
        accounts = accounts or {}
        plan_accounts = plan_accounts or {}
        account_rows: Dict[str, List[Tuple[int, int]]] = {}

        row = header + 1
        pen.text(_ref(2, row), "INDTÆGTER", bold=True, size=10, color=ORANGE, align="left")
        row += 1
        income_first = row
        for c in plan.income:
            self._write_plan_row(pen, row, c, targets, accounts, plan_accounts, account_rows, 1)
            row += 1
        income_last = row - 1
        row += 1
        income_total_row = row
        pen.text(_ref(2, row), "INDTÆGTER I ALT", bold=True, size=10, color=NAVY, align="left")
        if income_last >= income_first:
            pen.formula(_ref(4, row), f"=SUM({_range(4, income_first, 4, income_last)})",
                       fmt=FMT_CURRENCY, bold=True, color=NAVY, align="right")
        else:
            pen.number(_ref(4, row), 0, fmt=FMT_CURRENCY, bold=True, color=NAVY, align="right")
        row += 2

        expense_first = row
        for group in (GROUP_FIXED, GROUP_VARIABLE, GROUP_PERIODIC):
            members = plan.groups.get(group, [])
            if not members:
                continue
            pen.text(_ref(2, row), GROUP_LABELS[group], bold=True, size=10, color=ORANGE, align="left")
            row += 1
            for c in members:
                self._write_plan_row(pen, row, c, targets, accounts, plan_accounts, account_rows, -1)
                row += 1
            row += 1
        expense_last = row - 1

        expense_total_row = row
        pen.text(_ref(2, row), "UDGIFTER I ALT", bold=True, size=10, color=NAVY, align="left")
        pen.formula(_ref(4, row), f"=SUM({_range(4, expense_first, 4, expense_last)})",
                   fmt=FMT_CURRENCY, bold=True, color=NAVY, align="right")
        row += 2

        savings_row = row
        pen.text(_ref(2, row), "TIL OPSPARING", bold=True, size=11, color=ORANGE, align="left")
        pen.formula(_ref(4, row), f"={_ref(4, income_total_row)}-{_ref(4, expense_total_row)}",
                   fmt=FMT_SIGNED, bold=True, color=ORANGE, align="right")

        pen.freeze(_ref(1, header + 1))
        return {"income": income_total_row, "expense": expense_total_row,
               "savings": savings_row, "accounts": account_rows}

    def _write_plan_row(self, pen: Pen, row: int, c, targets, accounts, plan_accounts,
                        account_rows, sign: int) -> None:
        pen.text(_ref(2, row), c.category, size=10, color=DARK, align="left")
        pen.number(_ref(3, row), c.mean_all, fmt=FMT_CURRENCY, size=10, color=MUTED, align="right")

        key = (c.kind, c.category)
        target = targets.get(key, c.suggestion)
        pen.number(_ref(4, row), target, fmt=FMT_CURRENCY, size=10, bold=True, color=DARK,
                  align="right", bg=PEACH)
        pen.formula(_ref(5, row), f"={_ref(4, row)}-{_ref(3, row)}", fmt=FMT_SIGNED, size=10,
                   color=TEXT_GREY, align="right")

        account = plan_accounts.get(key) or accounts.get(c.category, "")
        pen.text(_ref(6, row), account, size=10, color=DARK, align="left")
        if account:
            account_rows.setdefault(account, []).append((row, sign))

        if c.sample:
            lines = [f"Gennemsnit: {c.mean_all:,.0f} kr./md".replace(",", "."),
                    f"{c.transaction_count} postering(er) i perioden."]
            pen.note(_ref(2, row), "\n".join(lines))

    def read_targets(self) -> Dict[Tuple[str, str], float]:
        if not self.has_sheet(SHEET_PLAN):
            return {}
        ws = self.wb[SHEET_PLAN]
        last_row = self.used_row_count(ws)
        out: Dict[Tuple[str, str], float] = {}
        labels = {fold("INDTÆGTER"), fold("INDTÆGTER I ALT"), fold("UDGIFTER I ALT"),
                 fold("TIL OPSPARING")} | {fold(GROUP_LABELS[g]) for g in GROUP_LABELS}
        for row in range(PLAN_HEADER_ROW + 1, last_row + 1):
            cat = str(ws.cell(row, 2).value or "").strip()
            if not cat or fold(cat) in labels:
                continue
            val = ws.cell(row, 4).value
            if not isinstance(val, (int, float)):
                continue
            out[(KIND_EXPENSE, cat)] = float(val)
            out[(KIND_INCOME, cat)] = float(val)
        return out

    def read_plan_accounts(self) -> Dict[Tuple[str, str], str]:
        if not self.has_sheet(SHEET_PLAN):
            return {}
        ws = self.wb[SHEET_PLAN]
        last_row = self.used_row_count(ws)
        out: Dict[Tuple[str, str], str] = {}
        labels = {fold("INDTÆGTER"), fold("INDTÆGTER I ALT"), fold("UDGIFTER I ALT"),
                 fold("TIL OPSPARING")} | {fold(GROUP_LABELS[g]) for g in GROUP_LABELS}
        for row in range(PLAN_HEADER_ROW + 1, last_row + 1):
            cat = str(ws.cell(row, 2).value or "").strip()
            if not cat or fold(cat) in labels:
                continue
            account = str(ws.cell(row, 6).value or "").strip()
            if account:
                out[(KIND_EXPENSE, cat)] = account
                out[(KIND_INCOME, cat)] = account
        return out

    # -- Prognose -----------------------------------------------------------
    def write_forecast(self, ws: Worksheet, plan_rows: Dict[str, object],
                       last_month_key: str, fallback_start: float) -> None:
        pen = self.pen(ws)
        for col, w in {1: 4, 2: 16, 3: 15, 4: 15}.items():
            pen.column_width(_col(col), w)

        pen.merge("B2:D2")
        pen.text("B2", "Prognose", font=FONT_TITLE, size=18, bold=True, color=ORANGE, align="left")
        pen.merge("B3:D3")
        pen.text("B3", f"Sådan udvikler saldoen sig de næste {FORECAST_MONTHS} måneder, hvis "
                       "du rammer målene i \"Budgetforslag\".",
                 size=9, italic=True, color=MUTED, align="left", wrap=True, valign="center")
        pen.row_height(3, 30)

        accounts_total = self.accounts_total()
        pen.text("B5", "Startsaldo", bold=True, size=10, color=NAVY, align="left")
        if accounts_total:
            pen.formula("D5", f"=SUM('{SHEET_ACCOUNTS}'!"
                              f"{_range(ACC_COL_BALANCE, ACC_FIRST_ROW, ACC_COL_BALANCE, ACC_FIRST_ROW + 20)})",
                       fmt=FMT_CURRENCY, size=10, color=DARK, align="right")
        else:
            pen.number("D5", fallback_start, fmt=FMT_CURRENCY, size=10, color=DARK, align="right", bg=PEACH)
        pen.text("B6", "Netto pr. måned med dine mål", size=10, color=TEXT_GREY, align="left")
        pen.formula("D6", f"='{SHEET_PLAN}'!{_ref(4, plan_rows['savings'])}",
                   fmt=FMT_SIGNED, size=10, bold=True, color=ORANGE, align="right")

        head = FORECAST_HEADER_ROW
        pen.text(_ref(2, head), "Måned", bold=True, size=11, color=NAVY, align="left")
        pen.text(_ref(3, head), "Netto/md", bold=True, size=10, color=NAVY, align="right")
        pen.text(_ref(4, head), "Saldo", bold=True, size=11, color=NAVY, align="right")

        start_month = _next_month(last_month_key)
        for i in range(FORECAST_MONTHS):
            row = head + 1 + i
            month = _add_months(start_month, i)
            pen.date(_ref(2, row), month, fmt=FMT_MONTH, size=10, color=DARK, align="left")
            pen.formula(_ref(3, row), "=$D$6", fmt=FMT_SIGNED, size=10, color=ORANGE, align="right")
            prev_ref = "$D$5" if i == 0 else _ref(4, row - 1)
            pen.formula(_ref(4, row), f"={prev_ref}+{_ref(3, row)}", fmt=FMT_CURRENCY, size=10,
                       bold=True, color=ORANGE, align="right")

        last_row = head + FORECAST_MONTHS
        self._add_forecast_chart(ws, head, last_row)
        pen.freeze(_ref(1, head))

    def _add_forecast_chart(self, ws: Worksheet, header_row: int, last_row: int) -> None:
        chart = LineChart()
        chart.title = "Saldoudvikling"
        chart.height = 8
        chart.width = 18
        data = Reference(ws, min_col=4, min_row=header_row, max_row=last_row)
        cats = Reference(ws, min_col=2, min_row=header_row + 1, max_row=last_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, "F2")

    def read_forecast_start(self) -> Optional[float]:
        if not self.has_sheet(SHEET_FORECAST):
            return None
        val = self.wb[SHEET_FORECAST]["D5"].value
        return float(val) if isinstance(val, (int, float)) else None

    # -- Alle måneder ---------------------------------------------------
    def write_months(self, ws: Worksheet, summary: Summary) -> None:
        pen = self.pen(ws)
        pen.column_width("A", 4)
        pen.column_width("B", 26)

        pen.merge("B2:D2")
        pen.text("B2", SHEET_MONTHS, font=FONT_TITLE, size=18, bold=True, color=ORANGE, align="left")

        head = 4
        pen.text(_ref(2, head), "Kategori", bold=True, size=11, color=NAVY, align="left")
        for col, month in enumerate(summary.months, start=3):
            pen.date(_ref(col, head), _month_start(month), fmt="mmm yyyy", bold=True, size=10,
                    color=NAVY, align="right")

        row = head + 1
        for cat in summary.expense_categories:
            pen.text(_ref(2, row), cat, size=10, color=DARK, align="left")
            for col, month in enumerate(summary.months, start=3):
                pen.number(_ref(col, row), summary.value(KIND_EXPENSE, cat, month),
                          fmt=FMT_CURRENCY, size=10, color=TEXT_GREY, align="right")
            row += 1
        row += 1
        for cat in summary.income_categories:
            pen.text(_ref(2, row), cat, size=10, color=DARK, align="left")
            for col, month in enumerate(summary.months, start=3):
                pen.number(_ref(col, row), summary.value(KIND_INCOME, cat, month),
                          fmt=FMT_CURRENCY, size=10, color=NAVY, align="right")
            row += 1

        pen.freeze(_ref(2, head + 1))

    # -- orchestration ----------------------------------------------------
    def build(self, transactions: Sequence[Transaction], ruleset: RuleSet,
             start_balance: float = 0.0) -> Summary:
        """Create every sheet from scratch."""
        summary = Summary(transactions)

        self.write_transactions(self.sheet(SHEET_TX), transactions)
        self.write_rules(self.sheet(SHEET_RULES), ruleset)
        self.ensure_accounts_sheet()

        if len(summary.months) > 1:
            plan = build_plan(summary, transactions)
            plan_rows = self.write_plan(self.sheet(SHEET_PLAN), plan)
            accounts_total = self.accounts_total()
            fallback = accounts_total if accounts_total else start_balance + summary.net
            self.write_forecast(self.sheet(SHEET_FORECAST), plan_rows, summary.months[-1], fallback)
            self.write_months(self.sheet(SHEET_MONTHS), summary)

        if "Sheet" in self.wb.sheetnames and len(self.wb.sheetnames) > 1:
            del self.wb["Sheet"]
        self._order_sheets()
        return summary

    def refresh(self, ruleset: RuleSet) -> Tuple[Optional[Summary], int]:
        """Re-apply the rules to every transaction and rebuild the summaries."""
        transactions = self.read_transactions()
        if not transactions:
            return None, 0
        ws = self.wb[SHEET_TX]
        account_numbers = self.read_account_numbers()
        changed = 0
        for t in transactions:
            if account_numbers and text_mentions_account(t.text, account_numbers):
                category = IGNORED_CATEGORY
            else:
                category = ruleset.categorise(t.text)
                if category == DEFAULT_CATEGORY and t.category:
                    category = t.category
            if category != t.category:
                col = COL_INC_CAT if t.income else COL_EXP_CAT
                ws.cell(t.source_row, col, category)
                t.category = category
                changed += 1
        self._apply_ignored_styling(ws, transactions)
        summary = self.rebuild_summaries(transactions)
        return summary, changed

    def _apply_ignored_styling(self, ws: Worksheet, transactions: Sequence[SheetTransaction]) -> None:
        pen = self.pen(ws)
        for t in transactions:
            first_col = COL_INC_DATE if t.income else COL_EXP_DATE
            row = t.source_row
            if t.category == IGNORED_CATEGORY:
                pen.style(_range(first_col, row, first_col + 3, row), color=MUTED, italic=True)
            else:
                pen.style(_range(first_col, row, first_col, row), color=MUTED, italic=False)
                pen.style(_range(first_col + 1, row, first_col + 3, row), color=TEXT_GREY, italic=False)

    def append_transactions(self, new_transactions: Sequence[Transaction]
                            ) -> Tuple[int, int, bool, Optional[Summary]]:
        """Add more transactions without touching what is already there.

        Returns (added_count, skipped_count, truncated, summary). ``summary``
        is None when there was nothing new to add.
        """
        from collections import Counter

        existing = self.read_transactions()
        existing_keys = Counter((t.date, t.text, round(t.amount, 2)) for t in existing)

        added: List[Transaction] = []
        skipped = 0
        for t in new_transactions:
            key = (t.date, t.text, round(t.amount, 2))
            if existing_keys[key] > 0:
                existing_keys[key] -= 1
                skipped += 1
            else:
                added.append(t)

        if not added:
            return 0, skipped, False, None

        ws = self.wb[SHEET_TX]
        existing_expenses = [t for t in existing if not t.income]
        existing_income = [t for t in existing if t.income]
        new_expenses = [t for t in added if t.amount < 0]
        new_income = [t for t in added if t.amount >= 0]

        room_exp = max(TX_MAX_ROW - TX_FIRST_ROW + 1 - len(existing_expenses), 0)
        room_inc = max(TX_MAX_ROW - TX_FIRST_ROW + 1 - len(existing_income), 0)
        truncated = False
        if len(new_expenses) > room_exp:
            new_expenses = new_expenses[:room_exp]
            truncated = True
        if len(new_income) > room_inc:
            new_income = new_income[:room_inc]
            truncated = True

        self._write_block(ws, new_expenses, COL_EXP_DATE, TX_FIRST_ROW + len(existing_expenses))
        self._write_block(ws, new_income, COL_INC_DATE, TX_FIRST_ROW + len(existing_income))

        new_exp_count = len(existing_expenses) + len(new_expenses)
        new_inc_count = len(existing_income) + len(new_income)
        last = TX_FIRST_ROW + max(new_exp_count, new_inc_count, 1) - 1
        pen = self.pen(ws)
        self._style_tx_range(pen, last)

        all_transactions = self.read_transactions()
        self._apply_ignored_styling(ws, all_transactions)
        self._add_category_dropdown(ws, COL_EXP_CAT, last)
        self._add_category_dropdown(ws, COL_INC_CAT, last)

        summary = self.rebuild_summaries(all_transactions)
        return len(added), skipped, truncated, summary

    def rebuild_summaries(self, transactions: Sequence[SheetTransaction]) -> Summary:
        """Recreate Budgetforslag/Prognose/Alle måneder from a transaction list."""
        summary = Summary(transactions)

        targets = self.read_targets()
        plan_accounts = self.read_plan_accounts()
        category_accounts = self.read_category_accounts()
        forecast_start = self.read_forecast_start()

        for name in (SHEET_PLAN, SHEET_FORECAST, SHEET_MONTHS):
            if self.has_sheet(name):
                del self.wb[name]

        self.ensure_accounts_sheet()

        if len(summary.months) > 1:
            plan = build_plan(summary, transactions)
            plan_rows = self.write_plan(self.sheet(SHEET_PLAN), plan, targets,
                                        category_accounts, plan_accounts)
            if forecast_start is not None:
                fallback = forecast_start
            else:
                accounts_total = self.accounts_total()
                fallback = accounts_total if accounts_total else summary.net
            self.write_forecast(self.sheet(SHEET_FORECAST), plan_rows, summary.months[-1], fallback)
            self.write_months(self.sheet(SHEET_MONTHS), summary)

        self._order_sheets()
        return summary

    def _order_sheets(self) -> None:
        order = [SHEET_PLAN, SHEET_FORECAST, SHEET_TX, SHEET_MONTHS, SHEET_RULES, SHEET_ACCOUNTS]
        for name in reversed(order):
            if name in self.wb.sheetnames:
                self.wb.move_sheet(name, offset=-self.wb.sheetnames.index(name))

    # -- persistence --------------------------------------------------------
    def save(self, path: str) -> None:
        self.wb.save(path)

    @classmethod
    def load(cls, path: str) -> "BudgetBook":
        return cls(load_workbook(path))
